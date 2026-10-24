"""Indexer: document loading, BM25 index, and vector (TF-IDF/embedding) index."""

import os
import re
import pickle
import hashlib
from html.parser import HTMLParser
from pathlib import Path
from typing import List, Dict, Any

import numpy as np

# Optional heavy deps
try:
    from rank_bm25 import BM25Plus
    HAS_BM25 = True
except ImportError:
    HAS_BM25 = False

try:
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
    HAS_SKLEARN = True
except ImportError:
    HAS_SKLEARN = False

HAS_ST = False


class _HTMLTextExtractor(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: List[str] = []
        self.ignored_tags: List[str] = []

    def handle_starttag(self, tag: str, attrs):
        if tag in {"script", "style"}:
            self.ignored_tags.append(tag)

    def handle_endtag(self, tag: str):
        if self.ignored_tags and tag == self.ignored_tags[-1]:
            self.ignored_tags.pop()

    def handle_data(self, data: str):
        if not self.ignored_tags:
            self.parts.append(data)


def _st_available() -> bool:
    global HAS_ST
    if HAS_ST:
        return True
    try:
        import sentence_transformers  # noqa: F401
        HAS_ST = True
        return True
    except ImportError:
        return False


SUPPORTED_EXT = {
    ".txt",
    ".md",
    ".rst",
    ".log",
    ".csv",
    ".json",
    ".html",
    ".htm",
    ".pdf",
    ".docx",
}
INDEX_FORMAT_VERSION = 3


def _read_file(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in {".txt", ".md", ".rst", ".log", ".csv", ".json"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext in {".html", ".htm"}:
        extractor = _HTMLTextExtractor()
        extractor.feed(path.read_text(encoding="utf-8", errors="ignore"))
        return re.sub(r"\s+", " ", " ".join(extractor.parts)).strip()
    if ext == ".pdf":
        try:
            from pdfminer.high_level import extract_text
        except ImportError as exc:
            raise RuntimeError(
                "PDF indexing requires pdfminer.six; install it with "
                "`python -m pip install pdfminer.six`."
            ) from exc
        return extract_text(str(path)) or ""
    if ext == ".docx":
        try:
            import docx
        except ImportError as exc:
            raise RuntimeError(
                "DOCX indexing requires python-docx; install it with "
                "`python -m pip install python-docx`."
            ) from exc
        document = docx.Document(str(path))
        content = [paragraph.text for paragraph in document.paragraphs]
        content.extend(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        return "\n".join(content)
    return ""


def tokenize(text: str) -> List[str]:
    return re.findall(r"[a-zA-Z0-9]+", (text or "").lower())


def load_documents(data_dir: str) -> List[Dict[str, Any]]:
    docs = []
    root = Path(data_dir)
    if not root.exists():
        return docs
    for p in sorted(root.rglob("*")):
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXT:
            text = _read_file(p)
            if not text.strip():
                continue
            docs.append({
                "id": hashlib.md5(str(p).encode()).hexdigest()[:10],
                "path": str(p),
                "title": p.name,
                "text": text,
            })
    return docs


class DualIndex:
    def __init__(self, data_dir: str = "data", index_dir: str = ".index", use_embeddings: bool = False):
        self.data_dir = data_dir
        self.index_dir = index_dir
        self.use_embeddings = use_embeddings and HAS_ST
        self.documents: List[Dict[str, Any]] = []
        self.bm25 = None
        self.bm25_tokens: List[List[str]] = []
        # vector side
        self.vectorizer = None
        self.vector_matrix = None  # (n_docs, n_features) or (n_docs, dim)
        self.embeddings = None
        self.st_model = None

    def build(self):
        self.documents = load_documents(self.data_dir)
        self.bm25 = None
        self.vectorizer = None
        self.vector_matrix = None
        self.embeddings = None
        self.st_model = None
        searchable_texts = [
            f'{document["title"]}\n{document["text"]}'
            for document in self.documents
        ]
        corpus_tokens = [tokenize(text) for text in searchable_texts]
        self.bm25_tokens = corpus_tokens
        if HAS_BM25 and corpus_tokens:
            self.bm25 = BM25Plus(corpus_tokens)
        # vector index
        texts = searchable_texts
        if self.use_embeddings and _st_available() and texts:
            from sentence_transformers import SentenceTransformer
            self.st_model = SentenceTransformer("all-MiniLM-L6-v2")
            self.embeddings = self.st_model.encode(texts, convert_to_numpy=True, normalize_embeddings=True)
        elif HAS_SKLEARN and texts:
            self.vectorizer = TfidfVectorizer(stop_words="english", max_features=5000)
            self.vector_matrix = self.vectorizer.fit_transform(texts)
        self._save()

    def _save(self):
        os.makedirs(self.index_dir, exist_ok=True)
        with open(os.path.join(self.index_dir, "index.pkl"), "wb") as f:
            pickle.dump({
                "format_version": INDEX_FORMAT_VERSION,
                "documents": self.documents,
                "bm25_tokens": self.bm25_tokens,
                "bm25": self.bm25,
                "vectorizer": self.vectorizer,
                "vector_matrix": self.vector_matrix,
                "embeddings": self.embeddings,
                "use_embeddings": self.use_embeddings,
            }, f)

    def load(self) -> bool:
        path = os.path.join(self.index_dir, "index.pkl")
        if not os.path.exists(path):
            return False
        with open(path, "rb") as f:
            data = pickle.load(f)
        if data.get("format_version") != INDEX_FORMAT_VERSION:
            return False
        self.documents = data["documents"]
        self.bm25_tokens = data["bm25_tokens"]
        self.bm25 = data["bm25"]
        self.vectorizer = data["vectorizer"]
        self.vector_matrix = data["vector_matrix"]
        self.embeddings = data["embeddings"]
        self.use_embeddings = data.get("use_embeddings", False)
        if self.use_embeddings and _st_available():
            try:
                from sentence_transformers import SentenceTransformer
                self.st_model = SentenceTransformer("all-MiniLM-L6-v2")
            except Exception:
                self.st_model = None
        return True

    def ensure_loaded(self):
        if not self.documents:
            if not self.load():
                self.build()

    @property
    def num_docs(self) -> int:
        return len(self.documents)
